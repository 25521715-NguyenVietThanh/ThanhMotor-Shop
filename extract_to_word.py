import ast
import os
from docx import Document
from docx.shared import RGBColor

def extract_docstrings_from_file(file_path):
    """
    Đọc file Python và trích xuất các docstring của Module, Class, Function.
    """
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            source = f.read()
        tree = ast.parse(source)
    except Exception:
        # Bỏ qua nếu file bị lỗi mã hóa hoặc không phải cú pháp Python chuẩn
        return []

    results = []

    # 1. Lấy Docstring của cả file (Module)
    module_doc = ast.get_docstring(tree)
    if module_doc:
        results.append(("Module", os.path.basename(file_path), module_doc))

    # 2. Lấy Docstring của các Class và Function bên trong
    for node in ast.walk(tree):
        if isinstance(node, ast.ClassDef):
            doc = ast.get_docstring(node)
            if doc:
                results.append(("Class", node.name, doc))

        elif isinstance(node, ast.FunctionDef):
            doc = ast.get_docstring(node)
            if doc:
                results.append(("Function", node.name, doc))

    return results


def extract_to_word(project_path):
    """
    Quét dự án và tạo file Word chứa toàn bộ Docstring.
    """
    doc = Document()
    
    # Thêm tiêu đề chính và căn giữa
    title = doc.add_heading("TÀI LIỆU DOCSTRING PROJECT (THANHMOTOR SHOP)", 0)
    title.alignment = 1

    for root, dirs, files in os.walk(project_path):
        # Thông minh hơn: Bỏ qua các thư mục không cần thiết để file Word sạch sẽ
        if any(ignored_dir in root for ignored_dir in ["venv", "__pycache__", ".git", "migrations"]):
            continue

        for file in files:
            if file.endswith(".py"):
                file_path = os.path.join(root, file)
                docs = extract_docstrings_from_file(file_path)

                if docs:
                    # Tiêu đề File (Heading 1)
                    doc.add_heading(f"File: {file}", level=1)

                    for doc_type, name, content in docs:
                        # Tiêu đề Class/Function (Heading 2)
                        doc.add_heading(f"{doc_type}: {name}", level=2)
                        
                        # Nội dung Docstring
                        p = doc.add_paragraph()
                        # Dùng .strip() để xóa các khoảng trắng/xuống dòng thừa thãi
                        run = p.add_run(content.strip())
                        
                        # Làm đẹp: In nghiêng và đổi sang màu xám cho chuyên nghiệp
                        run.italic = True
                        run.font.color.rgb = RGBColor(89, 89, 89)
                        
                        # Cố tình thêm một dòng trống để KHÔNG BAO GIỜ bị dính chữ vào nhau nữa
                        doc.add_paragraph()

    # Lưu file Word
    output_name = "Tai_lieu_docstring.docx"
    doc.save(output_name)
    print(f"✅ Đã trích xuất thành công! Hãy kiểm tra file: {output_name}")


if __name__ == "__main__":
    print("⏳ Đang quét toàn bộ file Python trong dự án...")
    # Quét từ thư mục hiện tại trở xuống
    extract_to_word(".")